"""Pipeline wrapper that runs the existing CLI pipeline inside a temp directory.

This module provides a thin adapter between the FastAPI web layer and the
existing ``src.main.process_connections`` function.  It writes the uploaded
CSV bytes to a temporary file, executes the full analysis pipeline (which
writes outputs to a temp output directory), then collects the generated
files into an in-memory ZIP archive and extracts KPI summary data from
the generated Excel workbook.

All temporary files are cleaned up unconditionally via ``try / finally``.
"""

from __future__ import annotations

import io
import logging
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Dict, Tuple

import openpyxl

from src.main import process_connections

logger = logging.getLogger("linkedin_analyzer")


def run_pipeline(
    csv_bytes: bytes,
    filename: str = "Connections.csv",
    *,
    role_filter: str | None = None,
    company_filter: str | None = None,
    company_contains: str | None = None,
    search_query: str | None = None,
    location_filter: str | None = None,
) -> Tuple[bytes, Dict[str, Any]]:
    """Execute the LinkedIn Network Analyzer pipeline on uploaded CSV data.

    Args:
        csv_bytes: Raw bytes of the uploaded CSV file.
        filename: Original filename (used for temp file naming).
        role_filter: Optional role filter (founder, cto, ceo, etc.).
        company_filter: Optional exact company name filter.
        company_contains: Optional company substring filter.
        search_query: Optional keyword search across name/company/position.
        location_filter: Optional location filter.

    Returns:
        Tuple of ``(zip_bytes, summary_dict)`` where:
        - ``zip_bytes`` is the ZIP archive containing all generated outputs.
        - ``summary_dict`` is a dictionary of KPI metrics extracted from the
          generated Excel workbook's Summary sheet.

    Raises:
        ValueError: If the CSV is invalid or the pipeline returns an error.
        Exception: Re-raises any unexpected pipeline error after cleanup.
    """
    tmp_dir = tempfile.mkdtemp(prefix="lna_web_")
    tmp_path = Path(tmp_dir)

    try:
        # Write uploaded CSV to temp file
        input_csv = tmp_path / filename
        input_csv.write_bytes(csv_bytes)

        output_dir = tmp_path / "output"
        output_dir.mkdir()

        log_dir = tmp_path / "logs"
        log_dir.mkdir()

        # Run the existing pipeline — returns 0 on success, non-zero on error
        exit_code = process_connections(
            input_path=str(input_csv),
            output_dir=str(output_dir),
            log_dir=str(log_dir),
            log_level="WARNING",  # reduce noise in web context
            role_filter=role_filter,
            company_filter=company_filter,
            company_contains=company_contains,
            search_query=search_query,
            location_filter=location_filter,
        )

        if exit_code != 0:
            raise ValueError(
                "Pipeline processing failed. The CSV file may not be a valid "
                "LinkedIn connections export, or it may contain no data rows."
            )

        # Collect generated output files into a ZIP archive
        zip_buffer = io.BytesIO()
        output_files = list(output_dir.iterdir())

        if not output_files:
            raise ValueError(
                "Pipeline completed but produced no output files. "
                "The CSV may not contain any senior contacts."
            )

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for fpath in output_files:
                if fpath.is_file():
                    zf.write(fpath, arcname=fpath.name)

        zip_bytes = zip_buffer.getvalue()

        # Extract KPI summary from the generated Excel workbook
        summary = _extract_summary(output_dir)

        return zip_bytes, summary

    finally:
        # Unconditional cleanup — delete all temp files
        _cleanup_temp_dir(tmp_path)


def _extract_summary(output_dir: Path) -> Dict[str, Any]:
    """Extract KPI metrics from the generated Excel workbook's Summary sheet.

    Reads the Summary tab (rows 4-12, columns A-B) of
    ``senior_connections.xlsx`` and builds a structured dictionary.

    Falls back to reading CSVs if the XLSX is unavailable.
    """
    xlsx_path = output_dir / "senior_connections.xlsx"
    summary: Dict[str, Any] = {}

    if xlsx_path.exists():
        try:
            wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
            if "Summary" in wb.sheetnames:
                ws = wb["Summary"]
                # KPI rows start at row 4: (Metric, Count/Value)
                for row in ws.iter_rows(min_row=4, max_row=12, min_col=1, max_col=2, values_only=True):
                    if row[0] and row[1] is not None:
                        key = str(row[0]).strip().lower().replace(" ", "_")
                        value = row[1]
                        # Convert numeric strings
                        if isinstance(value, str):
                            try:
                                value = int(value)
                            except ValueError:
                                try:
                                    value = float(value)
                                except ValueError:
                                    pass
                        summary[key] = value
                wb.close()
        except Exception as e:
            logger.warning(f"Could not extract summary from XLSX: {e}")

    # Fallback: count lines in CSVs if summary is empty
    if not summary:
        summary = _fallback_summary_from_csvs(output_dir)

    return summary


def _fallback_summary_from_csvs(output_dir: Path) -> Dict[str, Any]:
    """Build a summary from CSV line counts when XLSX parsing fails."""
    summary: Dict[str, Any] = {}
    csv_counts = {
        "senior_contacts": "senior_connections.csv",
        "founders": "founders.csv",
        "c-suite": "c_suite.csv",
        "directors": "directors.csv",
        "entrepreneurs": "entrepreneurs.csv",
    }
    for key, fname in csv_counts.items():
        fpath = output_dir / fname
        if fpath.exists():
            # Count data rows (subtract 1 for header)
            lines = fpath.read_text(encoding="utf-8-sig").strip().split("\n")
            summary[key] = max(len(lines) - 1, 0)
        else:
            summary[key] = 0

    return summary


def _cleanup_temp_dir(tmp_path: Path) -> None:
    """Recursively delete temp directory contents. Never raises."""
    import shutil
    try:
        if tmp_path.exists():
            shutil.rmtree(tmp_path, ignore_errors=True)
            logger.debug(f"Cleaned up temp directory: {tmp_path}")
    except Exception as e:
        logger.warning(f"Failed to clean up temp directory {tmp_path}: {e}")
